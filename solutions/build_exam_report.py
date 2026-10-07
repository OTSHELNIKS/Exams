from __future__ import annotations

from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    PageBreak, Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table as PDFTable,
    TableStyle, KeepTogether,
)

HERE = Path(__file__).resolve().parent
DOCX_PATH = HERE / "Отчет_решения_КИМ_09.02.07-5-2027.docx"
PDF_PATH = HERE / "Отчет_решения_КИМ_09.02.07-5-2027.pdf"
NAVY = "17365D"
BLUE = "4058D6"
PALE = "EAF0F8"
GRAY = "596579"

MONTH_DATA = [
    ("Январь", "413985", "5,34"),
    ("Февраль", "1278414", "16,48"),
    ("Март", "1208502", "15,57"),
    ("Апрель", "526072", "6,78"),
    ("Май", "279840", "3,61"),
    ("Июнь", "598556", "7,71"),
    ("Июль", "234172", "3,02"),
    ("Август", "332508", "4,29"),
    ("Сентябрь", "127949", "1,65"),
    ("Октябрь", "716678", "9,24"),
    ("Ноябрь", "1794540", "23,13"),
    ("Декабрь", "248464", "3,20"),
]

ENTITIES = [
    ("customers — Заказчики", "PK customer_id; name, INN, address, phone, party_type"),
    ("customer_orders — Заказы покупателей", "PK order_id; FK customer_id; order_date, status"),
    ("order_items — Позиции заказа", "PK order_item_id; FK order_id/product_id; quantity; снимок цены и скидка на единицу"),
    ("products — Товары", "PK product_id; source_product_id; name, category, is_placeholder"),
    ("resources — Материалы и операции", "PK resource_id; source_resource_id; name, resource_type"),
    ("product_specification — Спецификация", "PK specification_id; FK product_id/resource_id; quantity_per_product"),
    ("prices — Цены", "PK price_id; FK product_id ИЛИ resource_id; unit_price"),
    ("discounts — Скидки", "PK discount_id; FK product_id ИЛИ resource_id; discount_percent"),
    ("production_orders — Заказы на производство", "PK production_order_id; order_number, launch_date, department"),
    ("production_order_items — Строки выпуска", "PK production_order_item_id; FK production_order_id/product_id; quantity, unit"),
]

COST_ROWS = [
    ("Евровинт 6,5х5", "0,012", "432,083333", "10,37"),
    ("Мебельная деталь 500х800", "2", "37,50", "150,00"),
    ("Мебельная деталь 600х800", "4", "17,125", "137,00"),
    ("Опора", "4", "145,71875", "1165,75"),
    ("Столешница круглая", "1", "1242,50", "2485,00"),
    ("Распил ДСП, МДФ и листового материала", "1", "450,00", "900,00"),
    ("Сборка модулей", "1", "1400,00", "2800,00"),
    ("Упаковка", "1", "1900,00", "3800,00"),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_text(cell, text: str, bold: bool = False, color: str = "243247", size: float = 8.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.space_before = Pt(1)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_docx_table(doc: Document, headers: list[str], rows: list[tuple[str, ...]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, title in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], title, bold=True, color="FFFFFF", size=8.5)
        set_cell_shading(table.rows[0].cells[idx], NAVY)
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            set_cell_text(cells[idx], str(value), size=8.1)
            if len(table.rows) % 2 == 0:
                set_cell_shading(cells[idx], "F2F5FA")
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_docx_code(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.15)
    paragraph.paragraph_format.right_indent = Inches(0.15)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_together = True
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F3F6FA")
    p_pr.append(shading)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("243247")


def add_footer_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Стр. ")
    run.font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def create_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.32)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string("243247")
    normal.paragraph_format.space_after = Pt(5)
    for style_name, size, color in [("Heading 1", 17, NAVY), ("Heading 2", 12.5, BLUE), ("Heading 3", 10.5, NAVY)]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.text = "КИМ 09.02.07-5-2027  |  Отчёт о выполнении заданий"
    header.style = doc.styles["Normal"]
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    add_footer_page_number(section.footer.paragraphs[0])


    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(95)
    r = p.add_run("ОТЧЁТ")
    r.bold = True; r.font.name = "Calibri"; r.font.size = Pt(30); r.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("по решениям заданий демонстрационного экзамена")
    r.font.size = Pt(19); r.font.color.rgb = RGBColor.from_string(BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run("Специальность 09.02.07 «Информационные системы и программирование»\nКИМ 09.02.07-5-2027")
    r.font.size = Pt(12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(85)
    r = p.add_run("Выполнил(а): ______________________________\nДата: 07 октября 2026 г.")
    r.font.size = Pt(11)
    doc.add_page_break()

    doc.add_heading("Краткое резюме", 1)
    doc.add_paragraph(
        "Подготовлен комплект решений для семи заданий: ER-модель, реляционная база данных, "
        "SQL-расчёт себестоимости, приложение авторизации, API и тестовая коллекция, "
        "документация API, очищенный набор данных и временной ряд месячной выручки. "
        "Исходные экзаменационные файлы не изменялись; программные и табличные результаты находятся в папке solutions."
    )
    add_docx_table(doc, ["Задание", "Подготовленный результат"], [
        ("1", "ER-диаграмма в 3НФ и схема именования сущностей."),
        ("2", "SQLite-схема, импорт JSON, проверка внешних ключей."),
        ("3", "Запрос себестоимости заказа с учётом нормы расхода."),
        ("4", "Вход с капчей-пазлом, блокировкой, ролями и панелью администратора."),
        ("5", "REST API заметок и Postman/автоматические тесты."),
        ("6", "Заполненный шаблон документации API."),
        ("7", "Очистка 30 транзакций, годовые/месячные показатели, график со средним."),
    ], [0.7, 5.9])
    doc.add_heading("Состав рассмотренных материалов", 2)
    doc.add_paragraph(
        "Основной PDF КИМ содержит образцы заданий 1–6; отдельный PDF ТЗ_7 задаёт требования "
        "к очистке и анализу dataset_2026.xlsx. Приложение 1 описывает авторизацию и капчу, "
        "Приложение 2 — API заметок, Приложение 3 — шаблон документации. Архив «кимы.zip» содержит "
        "разложенные исходные файлы заданий 1–7; второй архив включает вложенные RAR-приложения."
    )

    doc.add_heading("Задание 1. Проектирование ER-диаграммы", 1)
    doc.add_paragraph(
        "Предметная область включает изготовление изделий из материалов и технологических операций, "
        "оформление заказов покупателей и производственных заказов. Справочник ресурсов объединяет "
        "материалы и операции, различаемые значением resource_type. Спецификация задаёт расход ресурса "
        "на одну единицу продукции."
    )
    add_docx_table(doc, ["Таблица", "Ключевые атрибуты и смысл"], ENTITIES, [1.65, 4.95])
    doc.add_heading("Связи и принципы нормализации", 2)
    for text in [
        "Заказчик 1:N заказы покупателей; заказ 1:N позиции заказа; товар 1:N позиции заказа.",
        "Товары и ресурсы связаны M:N через product_specification; поле quantity_per_product хранит норму расхода.",
        "Производственный заказ 1:N строк выпуска продукции. Цены и скидки относятся ровно к одному товару ИЛИ ресурсу; это обеспечено CHECK-ограничением.",
        "Справочники отделены от документов; снимок цены и скидка на единицу в позиции заказа не смешиваются с актуальным каталогом цен. Диаграмма приложена отдельным PDF: task1_er/ER_диаграмма.pdf.",
    ]:
        doc.add_paragraph(text, style="List Bullet")
    doc.add_heading("Выявленные проблемы исходных JSON", 2)
    doc.add_paragraph(
        "Для сохранения ссылочной целостности применены surrogate-ключи при повторяющихся исходных ID. "
        "В источнике повторяются product_id=887, resource_id=456, ID строк заказа 31 и 81, ID скидки 15. "
        "На отсутствующие в Товары.json ID 398, 881 и 942 созданы помеченные записи-заглушки. "
        "Цена с неоднозначной ссылкой product_id=887 назначена первой записи с таким ID; это явно отмечено в README базы данных."
    )

    doc.add_heading("Задание 2. Создание и заполнение базы данных", 1)
    doc.add_paragraph(
        "Создана схема SQLite с таблицами customers, customer_orders, order_items, products, resources, "
        "product_specification, prices, discounts, production_orders и production_order_items. "
        "Скрипт seed.py загружает предоставленные JSON-файлы, включая обязательный Заказчики.json, "
        "и после импорта запускает PRAGMA foreign_key_check. Нарушений внешних ключей не обнаружено."
    )
    add_docx_table(doc, ["Таблица", "Записей"], [
        ("customers", "6"), ("products", "132 (включая 3 заглушки)"),
        ("resources", "103"), ("product_specification", "74"),
        ("customer_orders", "70"), ("order_items", "70"),
        ("prices", "60"), ("discounts", "51"),
    ], [2.4, 4.2])
    add_docx_code(doc, "python3 solutions/task2_database/seed.py")
    doc.add_paragraph("DDL: task2_database/schema.sql. Готовая база: task2_database/manufacturing.sqlite3.")

    doc.add_heading("Задание 3. Запрос себестоимости заказа", 1)
    doc.add_paragraph(
        "Для каждого заказанного изделия запрос соединяет строку клиентского заказа со спецификацией "
        "и ценами ресурсов. Формула компонента: количество заказа × норма расхода × цена ресурса. "
        "Учитываются как материалы, так и технологические операции; скидка покупателю и цена реализации "
        "готового изделия в производственную себестоимость не включаются."
    )
    add_docx_code(doc, """WITH components AS (
  SELECT o.order_id, oi.quantity AS ordered_quantity,
         ps.specification_id, ps.quantity_per_product,
         p.price_id, p.unit_price
  FROM customer_orders o
  JOIN order_items oi ON oi.order_id = o.order_id
  LEFT JOIN product_specification ps ON ps.product_id = oi.product_id
  LEFT JOIN prices p ON p.resource_id = ps.resource_id AND p.product_id IS NULL
  WHERE o.order_id = :order_id
)
SELECT order_id,
       CASE WHEN COUNT(specification_id) = 0
                  OR COUNT(price_id) <> COUNT(specification_id)
            THEN NULL
            ELSE ROUND(SUM(ordered_quantity * quantity_per_product * unit_price), 2)
       END AS total_manufacturing_cost
FROM components GROUP BY order_id;""")
    doc.add_paragraph("Контрольный пример — заказ № 1: стол кухонный «Самобранка», количество 2.")
    add_docx_table(doc, ["Компонент", "Стоимость заказа, руб."], [
        (name, amount) for name, _, _, amount in COST_ROWS
    ] + [("Итого", "11448,12")], [4.6, 2.0])
    doc.add_paragraph("Полный запрос и построчная детализация: task3_query/order_cost.sql; результат совпадает с Расчет стоимости.xlsx.")

    doc.add_page_break()
    doc.add_heading("Задание 4. Модуль авторизации", 1)
    doc.add_paragraph(
        "Реализовано веб-приложение на Python стандартной библиотеки с SQLite. Форма требует логин и пароль; "
        "пароли хранятся как PBKDF2-HMAC-SHA256 с солью. После трёх последовательных ошибок капчи или пароля "
        "аккаунт блокируется. Капча построена как пазл 3×3 из приложенных изображений: фрагменты можно менять "
        "кликами или перетаскиванием, ответ проверяется сервером."
    )
    add_docx_table(doc, ["Функция", "Реализация"], [
        ("Роли", "Администратор и Пользователь."),
        ("Вход", "Обязательные login/password; сообщения об ошибке, блокировке и успешном входе."),
        ("Блокировка", "Три подряд неверных пароля/капчи; счётчик сбрасывается после успешного входа."),
        ("Администрирование", "Добавление пользователя, проверка уникальности логина, изменение логина/роли/пароля, снятие блокировки."),
        ("Хранилище", "SQLite; сессионная cookie HttpOnly/SameSite; пароль не хранится открытым текстом."),
    ], [1.55, 5.05])
    add_docx_code(doc, "python3 solutions/task4_auth_api/app.py\n# открыть http://127.0.0.1:8000\n# демонстрационные входы: admin / Admin123!; user / User123!")
    doc.add_paragraph("Реализация предназначена для экзаменационной демонстрации. Для промышленной эксплуатации потребуются HTTPS, CSRF-защита и постоянное хранилище сессий.")

    doc.add_heading("Задание 5. API и тестирование", 1)
    doc.add_paragraph(
        "Маршрут GET /notes (алиас GET /api/notes) возвращает заметки из таблицы notes с login автора из users. "
        "Поле title_user формируется как title + « - » + login; content не изменяется; дата выводится как ДД.ММ.ГГГГ. "
        "Поддерживаются необязательные параметры limit (1–1000) и user_id (>0)."
    )
    add_docx_code(doc, """[
  {
    "id": 5,
    "title_user": "Итоги - admin",
    "content": "Сохранить результаты тестирования.",
    "formatted_date": "05.10.2026"
  }
]""")
    add_docx_table(doc, ["Ситуация", "HTTP", "Ответ"], [
        ("Штатный запрос / нет подходящих заметок", "200", "JSON-массив / []"),
        ("Неверный/повторный параметр", "400", "JSON с полем error"),
        ("Ошибка базы данных", "500", "JSON: {\"error\":\"Ошибка подключения к базе данных\"}"),
        ("Метод, отличный от GET", "405", "JSON; заголовок Allow: GET"),
    ], [2.15, 0.65, 3.8])
    doc.add_paragraph("Коллекция Postman — task5_api/Exam_API.postman_collection.json. Запуск автоматических проверок:")
    add_docx_code(doc, "python3 -m unittest discover -s solutions/task4_auth_api/tests -v")
    doc.add_paragraph("Проверены 200, 400, JSON-формат, пустой массив, тестовый 500, 405, блокировка и административное добавление/разблокировка. Результат: 7 тестов пройдены.")

    doc.add_heading("Задание 6. Документация API", 1)
    doc.add_paragraph(
        "По предоставленному шаблону заполнены базовый URL, маршрут и алиас, параметры, описание обработки данных, "
        "формат JSON, статусы HTTP и примеры запросов/ответов для 200, 400 и 500. Готовый документ: "
        "task6_docs/API_документация.docx; скрипт сборки: task6_docs/build_docx.py."
    )

    doc.add_page_break()
    doc.add_heading("Задание 7. Очистка данных и анализ временного ряда", 1)
    doc.add_paragraph(
        "Исходный файл содержит 30 транзакций за 2026 год. Все строки сохранены; полных дубликатов нет, "
        "удаление строк не выполнялось. Приведены к единообразию ID транзакций, названия товаров и менеджеров, "
        "категории и числовые типы. Количества вида «3 шт.» и цены вида «294980 руб.» преобразованы в числа. "
        "Опечатка Dison исправлена на Dyson; категория нормализуется по товару, пропуск категории заполнен. "
        "В журнале очистки: 47 изменённых значений в 27 строках."
    )
    doc.add_heading("Месячная выручка и доля годовой выручки", 2)
    add_docx_table(doc, ["Месяц", "Выручка, руб.", "Доля, %"], MONTH_DATA + [("Итого", "7759680", "100,00")], [2.1, 2.2, 2.0])
    add_docx_table(doc, ["Показатель", "Результат"], [
        ("Среднее арифметическое за 12 месяцев", "646640 руб."),
        ("Максимальная выручка", "Ноябрь — 1794540 руб."),
        ("Минимальная выручка", "Сентябрь — 127949 руб."),
        ("Размах между максимумом и минимумом", "1666591 руб."),
        ("Линия среднего на графике задания 7.9", "646640 руб."),
        ("Интерпретация", "Ряд волатилен; устойчивого ежемесячного роста нет."),
    ], [3.4, 3.2])
    doc.add_paragraph(
        "График построен как линейный: месяцы расположены от января к декабрю; вторая серия — горизонтальная линия "
        "среднего значения. Файлы: task7_analysis/clean_dataset_2026.xlsx и analysis_2026.xlsx."
    )
    doc.add_heading("Ограничения и контроль", 2)
    doc.add_paragraph(
        "Транзакция TXN-0001 содержит цену 897270 руб./шт. за утюг Philips. Значение выглядит как сильный "
        "ценовой выброс, но в приложенных материалах нет проверенной цены или правила замены; поэтому исходное "
        "число сохранено и отмечено для сверки с первичным источником. Оно формирует ноябрьский пик."
    )
    doc.add_paragraph(
        "В загруженных файлах отсутствуют приложение Analytics_autotest и отдельные формулировки вопросов 7.2–7.9. "
        "Следовательно, официальный контрольный код задания 7.1 надо получить запуском экзаменационного приложения; "
        "технический SHA-256 файла не является его контрольной суммой."
    )

    doc.add_heading("Итоговый перечень файлов", 1)
    add_docx_table(doc, ["Задание", "Основные файлы решения"], [
        ("1", "task1_er/ER_диаграмма.pdf; ER-модель.mmd"),
        ("2", "task2_database/schema.sql; seed.py; manufacturing.sqlite3"),
        ("3", "task3_query/order_cost.sql; пример_расчёта.md"),
        ("4–5", "task4_auth_api/app.py; tests; task5_api/Exam_API.postman_collection.json"),
        ("6", "task6_docs/API_документация.docx"),
        ("7", "task7_analysis/clean_dataset_2026.xlsx; analysis_2026.xlsx; ответы_к_заданию_7.md"),
    ], [0.9, 5.7])
    doc.add_paragraph("Подробные инструкции по запуску и пояснения к данным собраны в solutions/README_RU.md.")
    doc.core_properties.title = "Отчёт по решениям КИМ 09.02.07-5-2027"
    doc.core_properties.subject = "Решения заданий 1–7"
    doc.core_properties.author = ""
    doc.save(DOCX_PATH)


def pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="RuTitle", fontName="DejaVu-Bold", fontSize=26, leading=32,
                              textColor=colors.HexColor("#" + NAVY), alignment=TA_CENTER, spaceAfter=14))
    styles.add(ParagraphStyle(name="RuSubtitle", fontName="DejaVu", fontSize=14, leading=20,
                              textColor=colors.HexColor("#" + BLUE), alignment=TA_CENTER, spaceAfter=10))
    styles.add(ParagraphStyle(name="RuH1", fontName="DejaVu-Bold", fontSize=16, leading=20,
                              textColor=colors.HexColor("#" + NAVY), spaceBefore=8, spaceAfter=6, keepWithNext=True))
    styles.add(ParagraphStyle(name="RuH2", fontName="DejaVu-Bold", fontSize=11.5, leading=15,
                              textColor=colors.HexColor("#" + BLUE), spaceBefore=6, spaceAfter=4, keepWithNext=True))
    styles.add(ParagraphStyle(name="RuBody", fontName="DejaVu", fontSize=8.8, leading=12.4,
                              textColor=colors.HexColor("#243247"), spaceAfter=5))
    styles.add(ParagraphStyle(name="RuSmall", fontName="DejaVu", fontSize=7.3, leading=9,
                              textColor=colors.HexColor("#243247"), spaceAfter=2))
    styles.add(ParagraphStyle(name="RuTableHeader", fontName="DejaVu-Bold", fontSize=7.3, leading=9,
                              textColor=colors.white, spaceAfter=2))
    styles.add(ParagraphStyle(name="RuCode", fontName="DejaVuMono", fontSize=6.6, leading=8.4,
                              textColor=colors.HexColor("#243247"), backColor=colors.HexColor("#F3F6FA"),
                              borderPadding=5, spaceBefore=3, spaceAfter=7))
    return styles


def pdf_table(headers, rows, widths, font_size=7.3):
    styles = pdf_styles()
    data = [[Paragraph(escape(str(v)), styles["RuTableHeader"]) for v in headers]]
    for row in rows:
        data.append([Paragraph(escape(str(v)), styles["RuSmall"]) for v in row])
    table = PDFTable(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + NAVY)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#C9D3E2")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5FA")]),
    ]))
    return table


def create_pdf() -> None:
    font_regular = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    font_bold = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")
    font_mono = Path("/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf")
    pdfmetrics.registerFont(TTFont("DejaVu", str(font_regular)))
    pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(font_bold)))
    pdfmetrics.registerFont(TTFont("DejaVuMono", str(font_mono)))
    styles = pdf_styles()
    story = []
    story += [Spacer(1, 42*mm), Paragraph("ОТЧЁТ", styles["RuTitle"]),
              Paragraph("по решениям заданий демонстрационного экзамена", styles["RuSubtitle"]),
              Spacer(1, 7*mm), Paragraph("Специальность 09.02.07<br/>КИМ 09.02.07-5-2027", styles["RuBody"]),
              Spacer(1, 26*mm), Paragraph("Выполнил(а): ______________________________<br/>Дата: 07 октября 2026 г.", styles["RuBody"]),
              PageBreak()]

    def p(text): story.append(Paragraph(text, styles["RuBody"]))
    def h1(text): story.append(Paragraph(escape(text), styles["RuH1"]))
    def h2(text): story.append(Paragraph(escape(text), styles["RuH2"]))
    def code(text): story.append(Preformatted(text, styles["RuCode"], maxLineLength=96))

    h1("Краткое резюме")
    p("Подготовлены решения для семи заданий: ER-модель, SQLite-база, SQL-расчёт себестоимости, приложение авторизации, API заметок и тесты, документация, очищенные данные и анализ месячной выручки. Исходные экзаменационные файлы не изменялись.")
    story.append(pdf_table(["Задание", "Результат"], [
        ("1", "ER-диаграмма в 3НФ."), ("2", "Схема БД, импорт JSON, проверка ссылок."),
        ("3", "Расчёт себестоимости заказа по нормам и ценам."),
        ("4", "Авторизация, капча-пазл, роли, блокировка, панель администратора."),
        ("5", "API GET /notes, Postman-коллекция и автоматические тесты."),
        ("6", "Документация в шаблоне из приложения."),
        ("7", "Очистка 30 транзакций, годовой анализ и график со средним."),
    ], [20*mm, 150*mm]))

    h1("Задание 1. Проектирование ER-диаграммы")
    p("Модель описывает изготовление изделий и заказы покупателей. Материалы и технологические операции объединены в resources и различаются resource_type. Спецификация связывает товары с ресурсами и хранит норму расхода на изделие.")
    story.append(pdf_table(["Таблица", "Ключевые поля / назначение"], ENTITIES, [42*mm, 128*mm]))
    h2("Связи и 3НФ")
    p("Заказчик 1:N заказы; заказ 1:N позиции; товар 1:N позиции. Товары и ресурсы связаны M:N через product_specification. Производственный заказ 1:N строк выпуска. Цены и скидки имеют ровно одну целевую ссылку — на товар либо на ресурс. Снимок цены и скидка в позиции заказа отделены от текущего прайс-листа.")
    p("В источниках обнаружены повторные ID товара 887, ресурса 456, строк заказа 31/81 и скидки 15. Они сохранены с surrogate-ключами. Для отсутствующих product_id 398, 881 и 942 созданы помеченные заглушки. ER-диаграмма: task1_er/ER_диаграмма.pdf.")

    h1("Задание 2. Создание базы данных")
    p("Созданы таблицы customers, customer_orders, order_items, products, resources, product_specification, prices, discounts, production_orders и production_order_items. Скрипт seed.py импортирует JSON и проверяет внешние ключи: нарушений нет.")
    story.append(pdf_table(["Таблица", "Число записей"], [
        ("customers", "6"), ("products", "132 (включая 3 заглушки)"), ("resources", "103"),
        ("product_specification", "74"), ("customer_orders", "70"), ("order_items", "70"),
        ("prices", "60"), ("discounts", "51"),
    ], [78*mm, 92*mm]))
    code("python3 solutions/task2_database/seed.py")

    h1("Задание 3. Запрос себестоимости")
    p("Себестоимость компонента заказа = количество заказанных изделий × норма расхода × цена ресурса. Материалы и операции учитываются одинаково. При отсутствии цены итог помечается как неполный, а не маскируется нулём.")
    code("""WITH components AS (
  SELECT o.order_id, oi.quantity AS ordered_quantity,
         ps.specification_id, ps.quantity_per_product, p.price_id, p.unit_price
  FROM customer_orders o JOIN order_items oi ON oi.order_id=o.order_id
  LEFT JOIN product_specification ps ON ps.product_id=oi.product_id
  LEFT JOIN prices p ON p.resource_id=ps.resource_id AND p.product_id IS NULL
  WHERE o.order_id=:order_id
)
SELECT order_id,
 CASE WHEN COUNT(specification_id)=0 OR COUNT(price_id)<>COUNT(specification_id)
      THEN NULL ELSE ROUND(SUM(ordered_quantity*quantity_per_product*unit_price),2)
 END AS total_manufacturing_cost
FROM components GROUP BY order_id;""")
    story.append(pdf_table(["Ресурс", "На заказ, руб."], [(r[0], r[3]) for r in COST_ROWS] + [("Итого", "11448,12")], [115*mm, 55*mm]))
    p("Заказ № 1 — стол «Самобранка», 2 шт.; итог 11448,12 руб. Это совпадает с приложенным расчётом стоимости.")

    h1("Задание 4. Модуль авторизации")
    p("Приложение написано на Python и SQLite. Есть роли «Администратор» и «Пользователь», обязательные поля входа, пазл-капча 3×3 на основе приложенных рисунков, серверная проверка решения и блокировка после трёх последовательных ошибок капчи или пароля. Пароли хранятся как PBKDF2-HMAC-SHA256 с солью.")
    p("Панель /admin/users позволяет добавлять учётные записи, проверять уникальность логина, менять логин, роль и пароль, блокировать и разблокировать пользователя. Демонстрационные входы: admin / Admin123! и user / User123!.")
    code("python3 solutions/task4_auth_api/app.py  # http://127.0.0.1:8000")

    h1("Задание 5. API и тестирование")
    p("GET /notes (алиас GET /api/notes) возвращает JSON-массив заметок. title_user = title + « - » + login; content передаётся без изменений; created_at преобразуется в ДД.ММ.ГГГГ. Необязательные параметры: limit (1–1000) и user_id (>0).")
    code("""[
  {"id":5,"title_user":"Итоги - admin",
   "content":"Сохранить результаты тестирования.",
   "formatted_date":"05.10.2026"}
]""")
    story.append(pdf_table(["Ситуация", "Код", "Ответ"], [
        ("Штатно; нет заметок", "200", "JSON-массив; при отсутствии данных []"),
        ("Некорректный параметр", "400", "JSON с error"),
        ("Сбой БД", "500", "JSON с описанием ошибки"),
        ("Метод не GET", "405", "JSON и Allow: GET"),
    ], [60*mm, 22*mm, 88*mm]))
    p("Коллекция Postman приложена. Автотесты покрывают 200, 400, 500, JSON-формат, пустой массив, 405, блокировку и администрирование. Пройдено 7 тестов.")

    h1("Задание 6. Документация API")
    p("По шаблону из приложения описаны базовый URL, GET /notes и алиас, параметры, формат данных, HTTP-статусы и примеры запросов/ответов 200, 400, 500. Документ: task6_docs/API_документация.docx.")

    h1("Задание 7. Очистка и анализ данных")
    p("Обработано 30 транзакций за 2026 год; полных дубликатов нет, строки не удалялись. Нормализованы ID, названия, категории и имена менеджеров; текстовые количества/цены преобразованы в числа; исправлена опечатка Dison → Dyson. В журнале — 47 исправлений в 27 строках.")
    h2("Выручка по месяцам")
    story.append(pdf_table(["Месяц", "Выручка, руб.", "Доля, %"], MONTH_DATA + [("Итого", "7759680", "100,00")], [55*mm, 65*mm, 50*mm]))
    story.append(pdf_table(["Показатель", "Результат"], [
        ("Средняя месячная выручка", "646640 руб."),
        ("Максимум", "Ноябрь — 1794540 руб."),
        ("Минимум", "Сентябрь — 127949 руб."),
        ("Размах", "1666591 руб."),
        ("Линия среднего на графике", "646640 руб."),
    ], [75*mm, 95*mm]))
    p("Линейный график содержит месяцы в хронологическом порядке и горизонтальную серию среднего значения. Ряд волатилен; устойчивого непрерывного роста нет.")
    h2("Ограничения и оговорки")
    p("В TXN-0001 цена утюга Philips равна 897270 руб./шт. Это сильный ценовой выброс, но в приложениях нет подтверждённой цены или правила замены; исходное число сохранено и помечено для сверки. Оно формирует ноябрьский пик.")
    p("В загруженных материалах нет Analytics_autotest и отдельных формулировок вопросов 7.2–7.9. Официальный контрольный код 7.1 надо получить запуском экзаменационного автотестера; SHA-256 файла не является контрольной суммой приложения.")

    h1("Файлы решения")
    story.append(pdf_table(["Задания", "Основные файлы"], [
        ("1", "task1_er/ER_диаграмма.pdf; ER-модель.mmd"),
        ("2", "task2_database/schema.sql; seed.py; manufacturing.sqlite3"),
        ("3", "task3_query/order_cost.sql; пример_расчёта.md"),
        ("4–5", "task4_auth_api/app.py; tests; task5_api/Exam_API.postman_collection.json"),
        ("6", "task6_docs/API_документация.docx"),
        ("7", "task7_analysis/clean_dataset_2026.xlsx; analysis_2026.xlsx; ответы_к_заданию_7.md"),
    ], [25*mm, 145*mm]))
    p("Подробные команды запуска приведены в solutions/README_RU.md.")

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#D7DFEA"))
        canvas.line(20*mm, 15*mm, 190*mm, 15*mm)
        canvas.setFont("DejaVu", 7)
        canvas.setFillColor(colors.HexColor("#596579"))
        canvas.drawString(20*mm, 10*mm, "КИМ 09.02.07-5-2027 · Отчёт по заданиям 1–7")
        canvas.drawRightString(190*mm, 10*mm, str(doc.page))
        canvas.restoreState()

    pdf = SimpleDocTemplate(str(PDF_PATH), pagesize=A4, rightMargin=20*mm, leftMargin=20*mm,
                            topMargin=17*mm, bottomMargin=21*mm, title="Отчёт по решениям КИМ 09.02.07-5-2027")
    pdf.build(story, onFirstPage=on_page, onLaterPages=on_page)


def main() -> None:
    create_docx()
    create_pdf()
    print(DOCX_PATH)
    print(PDF_PATH)


if __name__ == "__main__":
    main()
