#!/usr/bin/env python3
"""Создаёт подробный DOCX/PDF отчёт по всем заданиям с полными исходными листингами."""
from __future__ import annotations

from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parent
DOCX_OUT = ROOT / "Полное_пошаговое_решение_КИМ_09.02.07-5-2027.docx"
PDF_OUT = ROOT / "Полное_пошаговое_решение_КИМ_09.02.07-5-2027.pdf"
NAVY = "17365D"
BLUE = "4058D6"
TEXT = "243247"

MONTHS = [
    ("Январь", "413985", "5,34"), ("Февраль", "1278414", "16,48"),
    ("Март", "1208502", "15,57"), ("Апрель", "526072", "6,78"),
    ("Май", "279840", "3,61"), ("Июнь", "598556", "7,71"),
    ("Июль", "234172", "3,02"), ("Август", "332508", "4,29"),
    ("Сентябрь", "127949", "1,65"), ("Октябрь", "716678", "9,24"),
    ("Ноябрь", "1794540", "23,13"), ("Декабрь", "248464", "3,20"),
]
COSTS = [
    ("Евровинт 6,5х5", "0,012", "432,083333", "10,37"),
    ("Мебельная деталь 500х800", "2", "37,50", "150,00"),
    ("Мебельная деталь 600х800", "4", "17,125", "137,00"),
    ("Опора", "4", "145,71875", "1165,75"),
    ("Столешница круглая", "1", "1242,50", "2485,00"),
    ("Распил ДСП, МДФ и листового материала", "1", "450,00", "900,00"),
    ("Сборка модулей", "1", "1400,00", "2800,00"),
    ("Упаковка", "1", "1900,00", "3800,00"),
]

TASKS = [
("Задание 1. Проектирование ER-диаграммы", [
"Шаг 1. Изучить формы заказа покупателя, заказа на производство, спецификацию, цены, расчёт стоимости и JSON-справочники. Выделить сущности заказчика, заказа, строки заказа, товара, ресурса, спецификации, цены, скидки, производственного заказа и его строк.",
"Шаг 2. Задать первичные ключи. Для документа — отдельный идентификатор; строки заказа имеют собственный PK и FK на заказ/товар. Клиентский ID хранится как текст, чтобы не терять ведущие нули.",
"Шаг 3. Установить связи: customer 1:N customer_orders; customer_orders 1:N order_items; products 1:N order_items; products M:N resources через product_specification; production_orders 1:N production_order_items.",
"Шаг 4. В resources объединить материалы и операции, различая их значением resource_type. В спецификации хранить количество ресурса на одну единицу изделия.",
"Шаг 5. Для цен/скидок предусмотреть два nullable FK, но CHECK разрешает заполнить только один: ссылка либо на product, либо на resource. Цена/скидка в строке заказа при необходимости является историческим снимком и не должна зависеть от текущего каталога.",
"Шаг 6. Обеспечить 3НФ: атрибуты заказчика не повторяются в заказе; имя ресурса не дублируется в BOM; себестоимость рассчитывается запросом и не хранится как дублируемое поле.",
"Шаг 7. Представить модель как ER_диаграмма.pdf; редактируемый исходник — ER-модель.mmd. Полный DDL отражён в schema.sql.",
]),
("Задание 2. Создание и заполнение базы данных", [
"Шаг 1. Создать SQLite-БД по schema.sql. Включить PRAGMA foreign_keys=ON, определить NOT NULL, CHECK, FK и индексы по FK.",
"Шаг 2. Запустить seed.py. Он читает Заказчики.json, Заказы.json, Позиции_заказов.json, Товары.json, Материалы.json, Спецификация.json, Цены.json и Скидки.json и загружает их одной транзакцией. Обязательный импорт Заказчики.json выполняется.",
"Шаг 3. Сопоставить JSON-поля с колонками, например addres из источника записать как address. Числовые ID преобразовать в INTEGER, а customer_id оставить TEXT.",
"Шаг 4. Обработать конфликтные ссылки без отключения FK: повторные source ID сохранить отдельными surrogate PK с полем source_*_id. Отсутствующие товары 398, 881, 942 представить явно помеченными заглушками.",
"Шаг 5. Проверить PRAGMA foreign_key_check; ожидаемый результат — пустой набор нарушений. Проверить количество строк в таблицах и сохранить manufacturing.sqlite3.",
"Команда: python3 solutions/task2_database/seed.py. Скрипт перезагружает данные идемпотентно при той же схеме.",
]),
("Задание 3. Запрос себестоимости заказа", [
"Шаг 1. Передать номер заказа как параметр :order_id. Получить заказанные товары и количество из customer_orders и order_items.",
"Шаг 2. По product_id найти спецификацию. Каждая строка задаёт норму ресурса на одну единицу товара.",
"Шаг 3. По resource_id подобрать цену ресурса, включая и материалы, и технологические операции. Цена готового продукта и скидка покупателю не являются производственными затратами.",
"Шаг 4. Для каждого ресурса вычислить quantity заказа × quantity_per_product × unit_price; просуммировать, округлить до 2 знаков.",
"Шаг 5. Сравнить количество позиций спецификации и позиций с найденной ценой. Если цена или BOM отсутствуют, запрос возвращает NULL, а не заниженный частичный итог.",
"Полный итоговый запрос и запрос детализации приложены в task3_query/order_cost.sql.",
]),
("Задание 4. Авторизация и информационная система", [
"Шаг 1. Создать SQLite-таблицу users с уникальным логином, ролью Администратор/Пользователь, солью и PBKDF2-HMAC-SHA256-хешем пароля, is_blocked и failed_attempts.",
"Шаг 2. На форме входа сделать обязательные поля логина и пароля. Капча — пазл 3×3 из приложенных изображений. Игрок переставляет фрагменты мышью или кликами; сервер проверяет одноразовый challenge и правильный порядок.",
"Шаг 3. При ошибке пазла или пароля увеличить общий счётчик последовательных ошибок. На третьей подряд ошибке заблокировать пользователя и показать установленное заданием сообщение. При успешном входе сбросить счётчик.",
"Шаг 4. Создать рабочий стол после входа и вывести заданное сообщение об успешной авторизации. Сессию хранить в HttpOnly/SameSite cookie.",
"Шаг 5. Закрыть административный маршрут от роли Пользователь. На панели администратора добавить/редактировать логин, роль, пароль и статус блокировки; при добавлении проверять уникальность логина.",
"Шаг 6. Запуск: python3 solutions/task4_auth_api/app.py. Демо-входы: admin / Admin123! и user / User123!. Для промышленной эксплуатации необходимы HTTPS, CSRF-защита и постоянное хранилище сессий.",
]),
("Задание 5. API заметок и тестовая коллекция", [
"Шаг 1. В той же SQLite-БД создать notes: id, title, content, user_id (FK users), created_at. При первом запуске добавить тестовые записи.",
"Шаг 2. Реализовать GET /notes и алиас GET /api/notes. Соединить notes с users: title_user = title + ' - ' + login; content передать без изменения; created_at форматировать ДД.ММ.ГГГГ.",
"Шаг 3. Поддержать limit от 1 до 1000 и положительный user_id. Неизвестные, повторные, пустые/неверно типизированные параметры возвращают JSON 400.",
"Шаг 4. Пустая выборка — корректный 200 с []. Сбой БД — 500 с JSON-полем error. Неподдерживаемый метод — 405, Allow: GET. Успех и ошибки API всегда имеют application/json.",
"Шаг 5. Импортировать Exam_API.postman_collection.json и Exam_API.postman_environment.json. Обычный сервер слушает 8000; тестовый сбой для 500 включается только флагами EXAM_TEST_MODE=1 и EXAM_FORCE_DB_ERROR=1 на отдельном экземпляре 8001.",
"Шаг 6. Запустить: python3 -m unittest discover -s solutions/task4_auth_api/tests -v. Проверяются 200, 400, 500, JSON, [], 405, блокировка и админские операции; результат — 7 тестов пройдены.",
]),
("Задание 6. Документация API", [
"Шаг 1. На базе исходного DOCX-шаблона заполнить базовый URL, GET /notes и алиас, параметры, назначение метода и формат ответа.",
"Шаг 2. Описать статусы 200, 400, 500 и фактически реализованный 405, включая смысл пустого массива при отсутствии заметок.",
"Шаг 3. Привести примеры успешного запроса и ответов на неверный параметр и сбой БД; указать формат даты и структуру JSON-объекта.",
"Готовый документ: task6_docs/API_документация.docx; воспроизводимый генератор — task6_docs/build_docx.py.",
]),
("Задание 7. Очистка данных и временной ряд", [
"Шаг 1. Прочитать dataset_2026.xlsx; сохранить исходный порядок и все 30 строк. Удалять разрешено только полные дубликаты; в данном наборе их нет.",
"Шаг 2. Стандартизировать transaction_id, пробелы/регистр товаров и менеджеров. Исправить известную опечатку Пылесос Dison на Пылесос Dyson.",
"Шаг 3. Преобразовать значения количества '3 шт.' в число 3, цены '294980 руб.' — в 294980. Проверить даты на принадлежность 2026 году.",
"Шаг 4. Нормализовать категории по типу товара; исправить латинские символы-двойники в категории; заполнить пропущенную категорию по названию продукта.",
"Шаг 5. Выручка транзакции = quantity × price. Суммировать все транзакции по календарным месяцам без фильтрации магазинов, менеджеров и категорий. Для долей использовать общую годовую выручку.",
"Шаг 6. Построить линейный график января–декабря и добавить горизонтальную линию среднего арифметического 12 месячных значений. Результат: clean_dataset_2026.xlsx; анализ/график: analysis_2026.xlsx.",
"Шаг 7. Контроль: 30 строк; удалено 0; полных дублей 0; журнал содержит 47 исправлений в 27 строках. Годовая выручка 7759680, среднее 646640, максимум Ноябрь 1794540, минимум Сентябрь 127949.",
]),
]

SOURCE_FILES = [
    ("Приложение A. Исходная ER-модель", "task1_er/ER-модель.mmd"),
    ("Приложение B. SQL-схема", "task2_database/schema.sql"),
    ("Приложение C. Загрузчик JSON/SQLite", "task2_database/seed.py"),
    ("Приложение D. SQL-запросы задания 3", "task3_query/order_cost.sql"),
    ("Приложение E. Приложение заданий 4–5", "task4_auth_api/app.py"),
    ("Приложение F. Автоматические тесты", "task4_auth_api/tests/test_app.py"),
    ("Приложение G. Коллекция Postman", "task5_api/Exam_API.postman_collection.json"),
    ("Приложение H. Окружение Postman", "task5_api/Exam_API.postman_environment.json"),
    ("Приложение I. Генератор документации", "task6_docs/build_docx.py"),
    ("Приложение J. Скрипт очистки и анализа", "task7_analysis/clean_data.py"),
]


def add_docx_shading(paragraph, fill="F2F5F9"):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def add_docx_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = value
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(8.4)
            run.font.color.rgb = RGBColor(255, 255, 255)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), NAVY)
        cell._tc.get_or_add_tcPr().append(shd)
    for n, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(8)
        if n % 2 == 0:
            for cell in cells:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "F2F5F9")
                cell._tc.get_or_add_tcPr().append(shd)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_docx_code(doc, text):
    lines = text.splitlines() or [""]
    for start in range(0, len(lines), 36):
        if start:
            doc.add_page_break()
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.1)
        p.paragraph_format.right_indent = Inches(0.1)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = Pt(8.2)
        add_docx_shading(p)
        run = p.add_run("\n".join(lines[start:start + 36]))
        run.font.name = "Consolas"
        run.font.size = Pt(7)
        run.font.color.rgb = RGBColor.from_string(TEXT)


def setup_docx_header(section):
    header = section.header.paragraphs[0]
    header.text = "КИМ 09.02.07-5-2027 · ПОЛНОЕ ПОШАГОВОЕ РЕШЕНИЕ"
    header.runs[0].font.size = Pt(7.5)
    header.runs[0].font.color.rgb = RGBColor.from_string("596579")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Стр. ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.63)
    sec.bottom_margin = Inches(0.63)
    sec.left_margin = Inches(0.68)
    sec.right_margin = Inches(0.68)
    setup_docx_header(sec)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9.2)
    normal.paragraph_format.space_after = Pt(4)
    for name, size, color in [("Heading 1", 16, NAVY), ("Heading 2", 11.5, BLUE)]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(7)
        style.paragraph_format.space_after = Pt(3)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(74)
    run = p.add_run("ПОЛНОЕ ПОШАГОВОЕ РЕШЕНИЕ")
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("заданий 1–7 демонстрационного экзамена\n09.02.07-5-2027")
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(65)
    p.add_run("Выполнил(а): __________________________\nДата: 07 октября 2026 г.")
    doc.add_page_break()

    doc.add_heading("Как пользоваться документом", 1)
    doc.add_paragraph("Сначала приведены пошаговые действия и контрольные ответы по каждому заданию. В приложениях включены полные текстовые листинги ключевых исходников: ER-описание, DDL, загрузчик БД, запрос, приложение, тесты, Postman-коллекция, генератор документации и очистка данных. Готовые бинарные файлы лежат рядом в каталогах solutions/task1_er … task7_analysis.")
    doc.add_paragraph("Для пересборки таблиц требуется openpyxl; для PDF/DOCX-генераторов — reportlab и python-docx. Само веб-приложение заданий 4–5 использует только стандартную библиотеку Python.")
    doc.add_heading("Контрольные результаты", 2)
    add_docx_table(doc, ["Показатель", "Результат"], [
        ("Себестоимость заказа покупателя № 1", "11448,12 руб."),
        ("Транзакции в задании 7", "30; удалено 0; полных дублей 0"),
        ("Годовая / средняя месячная выручка", "7759680 / 646640 руб."),
        ("Максимум / минимум", "Ноябрь — 1794540; Сентябрь — 127949 руб."),
        ("Автоматические тесты приложения", "7 пройдены"),
    ], [2.8, 3.8])

    for title, steps in TASKS:
        doc.add_page_break()
        doc.add_heading(title, 1)
        for step in steps:
            p = doc.add_paragraph(step)
            p.paragraph_format.left_indent = Inches(0.16)
            p.paragraph_format.space_after = Pt(4)
        if title.startswith("Задание 1"):
            doc.add_heading("Сущности модели", 2)
            add_docx_table(doc, ["Таблица", "Назначение / ключи"], [
                ("customers", "PK customer_id; реквизиты покупателя/поставщика"),
                ("customer_orders", "PK order_id; FK customer_id; дата и статус"),
                ("order_items", "PK order_item_id; FK order_id/product_id; количество; снимок цены/скидки"),
                ("products", "PK product_id; название, категория, признак заглушки"),
                ("resources", "PK resource_id; материал либо операция"),
                ("product_specification", "PK specification_id; FK product/resource; норма расхода"),
                ("prices / discounts", "ссылка FK ровно на товар ИЛИ ресурс"),
                ("production_orders / items", "шапка производственного заказа и строки выпуска"),
            ], [2.15, 4.45])
            doc.add_paragraph("ER-диаграмма в PDF: task1_er/ER_диаграмма.pdf. Редактируемое описание: task1_er/ER-модель.mmd.")
        elif title.startswith("Задание 2"):
            doc.add_heading("Проверка импорта", 2)
            add_docx_table(doc, ["Таблица", "Записей"], [
                ("customers", "6"), ("products", "132"), ("resources", "103"),
                ("product_specification", "74"), ("customer_orders", "70"),
                ("order_items", "70"), ("prices", "60"), ("discounts", "51"),
            ], [2.3, 4.3])
            doc.add_paragraph("Проблемы исходника обработаны прозрачно: product_id 398/881/942 отсутствуют в справочнике; product_id 887 повторяется; resource_id 456 повторяется; повторяются ID строк заказа 31/81 и скидки 15. Строки не удалены. Неоднозначная ссылка на товар 887 назначена первой записи, потому что JSON не позволяет определить вариант.")
        elif title.startswith("Задание 3"):
            doc.add_heading("Ресурсная калькуляция заказа № 1", 2)
            add_docx_table(doc, ["Ресурс", "Норма / изделие", "Цена", "На заказ (2 шт.)"], COSTS + [("Итого", "—", "—", "11448,12")], [2.6, 1.15, 1.25, 1.6])
            doc.add_paragraph("10,37 + 150,00 + 137,00 + 1165,75 + 2485,00 + 900,00 + 2800,00 + 3800,00 = 11448,12 руб.")
            doc.add_paragraph("Полный итоговый запрос и детализация: task3_query/order_cost.sql.")
        elif title.startswith("Задание 4"):
            add_docx_table(doc, ["Сценарий", "Ожидаемый результат"], [
                ("Верный вход и капча", "Успех, рабочий стол, сброс счётчика"),
                ("Неверный пароль/капча", "Сообщение об ошибке и увеличение счётчика"),
                ("Третья подряд ошибка", "Блокировка учётной записи"),
                ("Работа администратора", "Добавление/редактирование и снятие блокировки"),
            ], [2.1, 4.5])
            doc.add_paragraph("Демо-логины в учебной БД: admin / Admin123!; user / User123!. Пароли PBKDF2-хешируются с солью.")
        elif title.startswith("Задание 5"):
            doc.add_heading("HTTP-статусы и схема ответа", 2)
            add_docx_table(doc, ["Код", "Ситуация", "Ответ"], [
                ("200", "Успешно / пустая выборка", "JSON-массив / []"),
                ("400", "Некорректный параметр", "JSON с error"),
                ("500", "Ошибка БД", "JSON с error"),
                ("405", "Метод не GET", "JSON; Allow: GET"),
            ], [0.75, 2.7, 3.15])
            doc.add_paragraph("Пример: {\"id\":5, \"title_user\":\"Итоги - admin\", \"content\":\"Сохранить результаты тестирования.\", \"formatted_date\":\"05.10.2026\"}.")
            doc.add_paragraph("Postman collection и environment находятся в task5_api/. Автотесты: 7 passed.")
        elif title.startswith("Задание 6"):
            doc.add_paragraph("Заполненный шаблон: task6_docs/API_документация.docx. В отчёте описаны параметры limit и user_id, поля объекта ответа и примеры для HTTP 200/400/500.")
        elif title.startswith("Задание 7"):
            doc.add_heading("Помесячные ответы", 2)
            add_docx_table(doc, ["Месяц", "Выручка, руб.", "Доля, %"], MONTHS + [("ИТОГО", "7759680", "100,00")], [2.3, 2.15, 2.15])
            add_docx_table(doc, ["Метрика", "Ответ"], [
                ("Среднее за 12 месяцев", "646640"), ("Максимум", "Ноябрь — 1794540"),
                ("Минимум", "Сентябрь — 127949"), ("Размах", "1666591"),
                ("Средняя линия графика 7.9", "646640"),
            ], [3.1, 3.5])
            doc.add_paragraph("Ценовой выброс TXN-0001 (утюг Philips, 897270 руб./шт.) отмечен и сохранён без неподтверждённой замены. Из-за него ноябрь становится максимумом.")
            doc.add_paragraph("В исходниках отсутствуют Analytics_autotest и точные формулировки вопросов 7.2–7.9; официальный контрольный код задания 7.1 получить можно только запуском экзаменационного приложения. SHA-256 файла не подменяет его контрольную сумму.")

    doc.add_page_break()
    doc.add_heading("Полные исходные листинги решения", 1)
    doc.add_paragraph("Ниже приведён полный текст ключевых исходников. Готовые PDF/XLSX/DOCX/SQLite-файлы приложены отдельно в каталоге solutions.")
    for index, (title, relative) in enumerate(SOURCE_FILES):
        if index:
            doc.add_page_break()
        doc.add_heading(title, 1)
        doc.add_paragraph(f"Файл: solutions/{relative}")
        path = ROOT / relative
        if path.exists():
            add_docx_code(doc, path.read_text(encoding="utf-8"))
        else:
            doc.add_paragraph("Файл отсутствует в текущей рабочей копии.")

    doc.core_properties.title = "Полное пошаговое решение КИМ 09.02.07-5-2027"
    doc.core_properties.subject = "Решения заданий 1–7 с полными исходными листингами"
    doc.core_properties.author = ""
    doc.save(DOCX_OUT)


def pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="FullTitle", fontName="DV-Bold", fontSize=25, leading=31, alignment=TA_CENTER, textColor=colors.HexColor("#" + NAVY), spaceAfter=12))
    styles.add(ParagraphStyle(name="FullSub", fontName="DV", fontSize=14, leading=20, alignment=TA_CENTER, textColor=colors.HexColor("#" + BLUE), spaceAfter=8))
    styles.add(ParagraphStyle(name="FullH1", fontName="DV-Bold", fontSize=13.5, leading=17, textColor=colors.HexColor("#" + NAVY), spaceBefore=7, spaceAfter=4, keepWithNext=True))
    styles.add(ParagraphStyle(name="FullH2", fontName="DV-Bold", fontSize=9.8, leading=12, textColor=colors.HexColor("#" + BLUE), spaceBefore=5, spaceAfter=3, keepWithNext=True))
    styles.add(ParagraphStyle(name="FullBody", fontName="DV", fontSize=7.8, leading=10.4, textColor=colors.HexColor("#" + TEXT), spaceAfter=3))
    styles.add(ParagraphStyle(name="FullSmall", fontName="DV", fontSize=6.7, leading=8.1, textColor=colors.HexColor("#" + TEXT), spaceAfter=1))
    styles.add(ParagraphStyle(name="FullTH", fontName="DV-Bold", fontSize=6.7, leading=8.1, textColor=colors.white))
    styles.add(ParagraphStyle(name="FullCode", fontName="DV-Mono", fontSize=6.0, leading=6.5, textColor=colors.HexColor("#152235"), backColor=colors.HexColor("#F2F5F9"), borderPadding=3, spaceAfter=2))
    return styles


def pdf_table(headers, rows, widths):
    styles = pdf_styles()
    data = [[Paragraph(escape(str(v)), styles["FullTH"]) for v in headers]]
    data += [[Paragraph(escape(str(v)), styles["FullSmall"]) for v in row] for row in rows]
    table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + NAVY)),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 2.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5FA")]),
    ]))
    return table


def build_pdf():
    pdfmetrics.registerFont(TTFont("DV", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
    pdfmetrics.registerFont(TTFont("DV-Bold", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"))
    pdfmetrics.registerFont(TTFont("DV-Mono", "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"))
    styles = pdf_styles()
    story = [Spacer(1, 45*mm), Paragraph("ПОЛНОЕ ПОШАГОВОЕ РЕШЕНИЕ", styles["FullTitle"]),
             Paragraph("заданий 1–7 демонстрационного экзамена", styles["FullSub"]), Spacer(1, 8*mm),
             Paragraph("Специальность 09.02.07<br/>КИМ 09.02.07-5-2027", styles["FullBody"]), Spacer(1, 25*mm),
             Paragraph("Выполнил(а): __________________________<br/>Дата: 07 октября 2026 г.", styles["FullBody"]), PageBreak()]

    def body(text):
        story.append(Paragraph(text, styles["FullBody"]))

    def heading(text):
        story.append(Paragraph(escape(text), styles["FullH1"]))

    def subheading(text):
        story.append(Paragraph(escape(text), styles["FullH2"]))

    def code_block(text):
        lines = text.splitlines() or [""]
        for start in range(0, len(lines), 58):
            story.append(Preformatted("\n".join(lines[start:start + 58]), styles["FullCode"], maxLineLength=120))

    heading("Как пользоваться документом")
    body("Первая часть содержит подробные действия и результаты по каждому заданию. В приложениях приведены полные текстовые листинги DDL, загрузчика, запроса, веб-приложения, тестов, Postman-коллекции и очистки данных. Бинарные ER-PDF, SQLite, XLSX и DOCX находятся в каталогах solutions/task1_er … task7_analysis.")
    story.append(pdf_table(["Показатель", "Результат"], [
        ("Себестоимость заказа №1", "11448,12 руб."), ("Строк задания 7", "30; удалено 0; дублей 0"),
        ("Годовая / средняя месячная выручка", "7759680 / 646640 руб."),
        ("Максимум / минимум", "Ноябрь — 1794540 / Сентябрь — 127949 руб."), ("Автотесты", "7 пройдены"),
    ], [68*mm, 102*mm]))

    for title, steps in TASKS:
        heading(title)
        for step in steps:
            body(escape(step))
        if title.startswith("Задание 1"):
            subheading("Сущности модели")
            story.append(pdf_table(["Таблица", "Ключи и смысл"], [
                ("customers", "PK customer_id; контрагенты"), ("customer_orders", "PK order_id; FK customer_id"),
                ("order_items", "PK строки; FK заказ/товар; qty, snapshot price/discount"),
                ("products", "PK product_id; каталог и заглушки"), ("resources", "PK resource_id; материал/операция"),
                ("product_specification", "PK; FK товар/ресурс; норма"),
                ("prices / discounts", "ровно одна FK-ссылка на товар или ресурс"),
                ("production_orders / items", "производственный заказ и строки выпуска"),
            ], [50*mm, 120*mm]))
            body("Кардинальности: customer 1:N orders; order 1:N order_items; product 1:N order_items; product M:N resource через specification; production order 1:N production items.")
            body("В источниках повторяются product 887, resource 456, order-item IDs 31/81, discount ID 15. Товары 398/881/942 отсутствуют в справочнике. Дубли сохранены surrogate keys, отсутствующие товары помечены заглушками; FK не отключались.")
            body("Диаграмма: task1_er/ER_диаграмма.pdf; редактируемая модель: task1_er/ER-модель.mmd.")
        elif title.startswith("Задание 2"):
            story.append(pdf_table(["Таблица", "Строк"], [
                ("customers", "6"), ("products", "132"), ("resources", "103"),
                ("product_specification", "74"), ("customer_orders", "70"),
                ("order_items", "70"), ("prices", "60"), ("discounts", "51"),
            ], [85*mm, 85*mm]))
            body("После загрузки PRAGMA foreign_key_check не находит нарушений. Снимок цены клиентской позиции оставлен NULL, поскольку JSON позиции не передаёт цену; скидка по умолчанию 0.")
        elif title.startswith("Задание 3"):
            subheading("Проверочный расчёт заказа №1")
            story.append(pdf_table(["Ресурс", "Норма", "Цена", "На заказ"], COSTS + [("ИТОГО", "—", "—", "11448,12")], [68*mm, 24*mm, 35*mm, 43*mm]))
            body("10,37 + 150,00 + 137,00 + 1165,75 + 2485,00 + 900,00 + 2800,00 + 3800,00 = 11448,12 руб.")
        elif title.startswith("Задание 4"):
            story.append(pdf_table(["Сценарий", "Ожидаемый результат"], [
                ("Верный вход", "Рабочий стол; сообщение об успехе; сброс счётчика"),
                ("Неверная капча/пароль", "Ошибка и увеличение failed_attempts"),
                ("Третья ошибка", "Блокировка аккаунта"),
                ("Администратор", "Добавление/редактирование/разблокировка"),
            ], [55*mm, 115*mm]))
            body("Демо-учётные записи: admin / Admin123! и user / User123!. Смените их перед эксплуатацией; приложение учебное, без HTTPS/CSRF-защиты.")
        elif title.startswith("Задание 5"):
            story.append(pdf_table(["Код", "Сценарий", "Ответ"], [
                ("200", "Штатно / данных нет", "JSON-массив / []"),
                ("400", "Некорректный параметр", "JSON с error"),
                ("500", "Сбой БД", "JSON с error"),
                ("405", "Метод не GET", "JSON и Allow: GET"),
            ], [20*mm, 68*mm, 82*mm]))
            code_block('[{"id":5,"title_user":"Итоги - admin","content":"Сохранить результаты тестирования.","formatted_date":"05.10.2026"}]')
            body("Коллекция Postman и environment в task5_api/. При тестовом 500 запускается второй экземпляр с EXAM_TEST_MODE=1 и EXAM_FORCE_DB_ERROR=1.")
        elif title.startswith("Задание 6"):
            body("Документ task6_docs/API_документация.docx заполнен по исходному шаблону; генератор — task6_docs/build_docx.py.")
        elif title.startswith("Задание 7"):
            subheading("Месячная выручка и доли")
            story.append(pdf_table(["Месяц", "Выручка, руб.", "Доля, %"], MONTHS + [("Итого", "7759680", "100,00")], [56*mm, 60*mm, 54*mm]))
            story.append(pdf_table(["Контрольный ответ", "Значение"], [
                ("Средняя за 12 месяцев", "646640 руб."), ("Максимум", "Ноябрь — 1794540 руб."),
                ("Минимум", "Сентябрь — 127949 руб."), ("Размах", "1666591 руб."),
                ("Линия среднего графика 7.9", "646640 руб."),
            ], [80*mm, 90*mm]))
            body("Выручка волатильна, устойчивого непрерывного роста нет. У TXN-0001 цена утюга Philips равна 897270 руб./шт.; значение отмечено как выброс, но сохранено, поскольку исправленная цена и правило её подстановки в ТЗ отсутствуют. Этот ряд формирует ноябрьский пик.")
            body("В загруженных материалах отсутствуют Analytics_autotest и конкретные тексты вопросов 7.2–7.9. Официальный контрольный код 7.1 нужно получить запуском приложения; технический SHA-256 не является контрольной суммой тестера.")

    story.append(PageBreak())
    heading("Полные текстовые листинги")
    body("Следующие приложения включают полный исходный текст файлов проекта. Чтобы запустить решение, используйте соответствующие команды из solutions/README_RU.md.")
    for index, (title, relative) in enumerate(SOURCE_FILES):
        if index:
            story.append(PageBreak())
        heading(title)
        body("Файл: solutions/" + escape(relative))
        path = ROOT / relative
        if path.exists():
            code_block(path.read_text(encoding="utf-8"))
        else:
            body("Файл не найден в каталоге решения.")

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#D7DFEA"))
        canvas.line(20*mm, 15*mm, 190*mm, 15*mm)
        canvas.setFont("DV", 6.5)
        canvas.setFillColor(colors.HexColor("#596579"))
        canvas.drawString(20*mm, 10*mm, "КИМ 09.02.07-5-2027 · Полное решение заданий 1–7")
        canvas.drawRightString(190*mm, 10*mm, str(doc.page))
        canvas.restoreState()

    pdf = SimpleDocTemplate(str(PDF_OUT), pagesize=A4, rightMargin=20*mm, leftMargin=20*mm,
                            topMargin=15*mm, bottomMargin=20*mm,
                            title="Полное пошаговое решение КИМ 09.02.07-5-2027")
    pdf.build(story, onFirstPage=on_page, onLaterPages=on_page)


def main():
    build_docx()
    build_pdf()
    print(DOCX_OUT)
    print(PDF_OUT)

if __name__ == "__main__":
    main()
