from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "Прил_6_ОЗ_КИМ_09.02.07-5-2027.docx"
OUTPUT = Path(__file__).resolve().parent / "API_документация.docx"


def set_cell(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    for index, line in enumerate(text.split("\n")):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(line)
        run.bold = bold
        run.font.size = Pt(8.5)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_code(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.2)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(31, 45, 61)


def main() -> None:
    doc = Document(TEMPLATE)
    doc.core_properties.title = "Документация API заметок"
    doc.core_properties.subject = "Задание 6 — описание разработанного API"
    doc.core_properties.author = "Участник демонстрационного экзамена"

    base_table = doc.tables[0]
    base_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell(base_table.cell(0, 0), "Базовый URL", bold=True)
    base_table.cell(0, 1).merge(base_table.cell(0, 3))
    set_cell(base_table.cell(0, 1), "http://localhost:8000")
    set_cell(base_table.cell(1, 0), "Название метода", bold=True)
    set_cell(base_table.cell(1, 1), "Параметры", bold=True)
    set_cell(base_table.cell(1, 2), "Описание метода", bold=True)
    set_cell(base_table.cell(1, 3), "Формат ответа", bold=True)
    set_cell(base_table.cell(2, 0), "GET /notes\nGET /api/notes (алиас)")
    set_cell(base_table.cell(2, 1), "limit — необязательное целое число 1–1000;\nuser_id — необязательное положительное целое число.")
    set_cell(base_table.cell(2, 2), "Возвращает заметки из notes, присоединяя login автора из users. title_user = title + ' - ' + login; content без изменений; created_at преобразуется в ДД.ММ.ГГГГ. Неизвестные, повторные и некорректные параметры дают 400. Если заметок нет, возвращается [].")
    set_cell(base_table.cell(2, 3), "200: массив JSON-объектов {id, title_user, content, formatted_date}. UTF-8; Content-Type: application/json.")
    for row in base_table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    status_table = doc.tables[1]
    set_cell(status_table.cell(0, 0), "Код", bold=True)
    set_cell(status_table.cell(0, 1), "Описание", bold=True)
    status_values = [
        ("200 OK", "Запрос обработан. Тело — JSON-массив; при отсутствии подходящих заметок — пустой массив []."),
        ("400 Bad Request", "Некорректный, повторный или неизвестный query-параметр; тело содержит поле error."),
        ("500 Internal Server Error", "Внутренняя ошибка/сбой БД; JSON вида {\"error\": \"Ошибка подключения к базе данных\"}."),
    ]
    for index, values in enumerate(status_values, start=1):
        set_cell(status_table.cell(index, 0), values[0])
        set_cell(status_table.cell(index, 1), values[1])
    extra = status_table.add_row().cells
    set_cell(extra[0], "405 Method Not Allowed")
    set_cell(extra[1], "Метод, отличный от GET, не поддерживается; заголовок Allow: GET, тело ошибки — JSON.")

    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(10)
    run = heading.add_run("Примеры запросов и ответов")
    run.bold = True
    run.font.size = Pt(12)

    doc.add_paragraph("1. Успешное получение двух последних заметок:")
    add_code(doc, "GET http://localhost:8000/notes?limit=2\nAccept: application/json")
    add_code(doc, """HTTP/1.1 200 OK
Content-Type: application/json; charset=utf-8

[
  {
    "id": 5,
    "title_user": "Итоги - admin",
    "content": "Сохранить результаты тестирования.",
    "formatted_date": "05.10.2026"
  },
  {
    "id": 4,
    "title_user": "Проверка базы - user",
    "content": "Убедиться, что внешние ключи включены.",
    "formatted_date": "04.10.2026"
  }
]""")

    doc.add_paragraph("2. Некорректный тип параметра:")
    add_code(doc, "GET http://localhost:8000/notes?limit=abc")
    add_code(doc, """HTTP/1.1 400 Bad Request
Content-Type: application/json; charset=utf-8

{"error":"Параметры limit и user_id должны быть положительными целыми числами; limit — от 1 до 1000."}""")

    doc.add_paragraph("3. Внутренняя ошибка базы данных (например, база недоступна):")
    add_code(doc, "GET http://localhost:8000/notes")
    add_code(doc, """HTTP/1.1 500 Internal Server Error
Content-Type: application/json; charset=utf-8

{"error":"Ошибка подключения к базе данных"}""")

    doc.add_paragraph("Схема ответа одной заметки:")
    add_code(doc, """{
  "id": 1,
  "title_user": "Конференция ИТ - user25",
  "content": "Текст заметки",
  "formatted_date": "15.03.2027"
}""")
    doc.add_paragraph("Параметры фильтрации не меняют формат объекта ответа. Порядок по умолчанию — от новых заметок к старым.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
