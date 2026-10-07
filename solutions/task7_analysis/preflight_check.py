from __future__ import annotations

import argparse
import sys
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.chart import LineChart

from clean_data import infer_category, normalize_manager, normalize_product, normalize_transaction_id

HERE = Path(__file__).resolve().parent
REPO_ROOT = HERE.parents[1]
DEFAULT_SOURCE = HERE / "dataset_2026.xlsx"
if not DEFAULT_SOURCE.is_file():
    DEFAULT_SOURCE = REPO_ROOT / "dataset_2026.xlsx"
DEFAULT_CLEAN = HERE / "clean_dataset_2026.xlsx"
DEFAULT_ANALYSIS = HERE / "analysis_2026.xlsx"
HEADERS = ["transaction_id", "date", "product", "category", "quantity", "price", "shop", "manager"]
MONTHS = [
    "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
    "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь",
]
EXPECTED = {
    "Строк в исходном наборе": 30,
    "Полных дубликатов": 0,
    "Строк удалено": 0,
    "Годовая выручка, руб.": 7_759_680,
    "Средняя месячная выручка, руб.": 646_640,
    "Месяц максимальной выручки": "Ноябрь",
    "Максимальная месячная выручка, руб.": 1_794_540,
    "Месяц минимальной выручки": "Сентябрь",
    "Минимальная месячная выручка, руб.": 127_949,
    "Размах месячной выручки, руб.": 1_666_591,
    "Строк с зафиксированными исправлениями": 27,
    "Количество исправлений по полям": 47,
}


def dec(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        return None
    return result if result.is_finite() else None


def check(condition: bool, label: str, failures: list[str]) -> None:
    print(f"[{'OK' if condition else 'ОШИБКА'}] {label}")
    if not condition:
        failures.append(label)


def main() -> int:
    parser = argparse.ArgumentParser(description='Локальная предварительная проверка файлов задания 7.')
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--clean", type=Path, default=DEFAULT_CLEAN)
    parser.add_argument("--analysis", type=Path, default=DEFAULT_ANALYSIS)
    args = parser.parse_args()
    failures: list[str] = []

    for label, path in (("Исходный набор", args.source), ("Очищенный набор", args.clean), ("Книга анализа", args.analysis)):
        check(path.is_file(), f"{label}: файл существует ({path})", failures)
    if failures:
        print("Проверка остановлена: не найдены входные/выходные файлы.")
        return 1

    source_wb = load_workbook(args.source, data_only=True, read_only=True)
    clean_wb = load_workbook(args.clean, data_only=True, read_only=True)
    analysis_wb = load_workbook(args.analysis, data_only=True, read_only=False)
    try:
        source_rows_all = list(source_wb.active.iter_rows(values_only=True))
        source_headers = list(source_rows_all[0][:8]) if source_rows_all else []
        source_rows = [tuple(row[:8]) for row in source_rows_all[1:] if any(v is not None for v in row[:8])]
        clean_sheet = clean_wb.active
        clean_rows_all = list(clean_sheet.iter_rows(values_only=True))
        clean_headers = list(clean_rows_all[0][:8]) if clean_rows_all else []
        clean_rows = [tuple(row[:8]) for row in clean_rows_all[1:] if any(v is not None for v in row[:8])]

        check(source_headers == HEADERS, "Исходные заголовки совпадают с заданием", failures)
        check(clean_headers == HEADERS, "Очищенная книга сохраняет исходные 8 колонок", failures)
        check(len(source_rows) == EXPECTED["Строк в исходном наборе"], "В источнике ровно 30 транзакций", failures)
        check(len(clean_rows) == len(source_rows), "Число строк сохранено: ни одна транзакция не потеряна", failures)
        duplicate_count = len(clean_rows) - len(set(clean_rows))
        check(duplicate_count == 0, "Полных дубликатов после очистки нет", failures)

        ids: list[str] = []
        dates_ok = True
        numbers_ok = True
        canonical_ok = len(source_rows) == len(clean_rows)
        for raw, row in zip(source_rows, clean_rows):
            tx_id, tx_date, product, category, quantity, price, shop, manager = row
            ids.append(str(tx_id))
            dates_ok &= isinstance(tx_date, (datetime, date)) and tx_date.year == 2026
            q, p = dec(quantity), dec(price)
            numbers_ok &= q is not None and p is not None and q > 0 and q == q.to_integral_value() and p >= 0
            expected_product = normalize_product(raw[2])
            expected_manager = normalize_manager(raw[7])
            canonical_ok &= str(tx_id) == normalize_transaction_id(raw[0])
            canonical_ok &= str(product) == expected_product
            canonical_ok &= str(category) == infer_category(expected_product)
            canonical_ok &= str(shop) == " ".join(str(raw[6] or "").split())
            canonical_ok &= str(manager) == expected_manager
        check(len(ids) == len(set(ids)) and all(len(tx_id) == 8 for tx_id in ids), "ID транзакций нормализованы и уникальны", failures)
        check(dates_ok, "Все даты — настоящие даты 2026 года", failures)
        check(numbers_ok, "Количество и цена — числовые; количество целое и положительное", failures)
        check(canonical_ok, "Названия, категории, магазины и менеджеры приведены к единому виду", failures)

        monthly = [Decimal("0") for _ in range(12)]
        line_revenues: list[Decimal] = []
        for row in clean_rows:
            tx_date, quantity, price = row[1], dec(row[4]), dec(row[5])
            revenue = quantity * price
            line_revenues.append(revenue)
            monthly[tx_date.month - 1] += revenue
        total = sum(monthly, Decimal("0"))
        average = total / Decimal(12)
        max_index = max(range(12), key=lambda i: monthly[i])
        min_index = min(range(12), key=lambda i: monthly[i])
        check(total == Decimal("7759680"), f"Годовая выручка рассчитана верно: {total}", failures)

        required_sheets = {"Данные", "Месячная выручка", "Итоги", "Журнал очистки"}
        check(required_sheets.issubset(set(analysis_wb.sheetnames)), "В книге анализа есть листы данных, месяцев, итогов и журнала", failures)
        if required_sheets.issubset(set(analysis_wb.sheetnames)):
            data_sheet = analysis_wb["Данные"]
            data_values = list(data_sheet.iter_rows(min_row=2, values_only=True))
            data_ok = len(data_values) == len(clean_rows)
            for row, clean, expected_revenue in zip(data_values, clean_rows, line_revenues):
                data_ok &= str(row[0]) == str(clean[0])
                actual = dec(row[8]) if len(row) > 8 else None
                data_ok &= actual == expected_revenue
            check(data_ok, "Выручка каждой строки равна quantity × price", failures)

            month_sheet = analysis_wb["Месячная выручка"]
            month_names = [month_sheet.cell(row, 1).value for row in range(2, 14)]
            actual_months = [dec(month_sheet.cell(row, 2).value) for row in range(2, 14)]
            actual_shares = [dec(month_sheet.cell(row, 3).value) for row in range(2, 14)]
            actual_average = [dec(month_sheet.cell(row, 6).value) for row in range(2, 14)]
            months_ok = month_names == MONTHS and actual_months == monthly
            expected_shares = [value / total * Decimal("100") for value in monthly]
            shares_ok = (
                all(value is not None for value in actual_shares)
                and all(abs(actual - expected) <= Decimal("0.01") for actual, expected in zip(actual_shares, expected_shares))
                and abs(sum(actual_shares, Decimal("0")) - Decimal("100")) <= Decimal("0.02")
            )
            avg_ok = all(value is not None and abs(value - average) < Decimal("0.01") for value in actual_average)
            check(months_ok, "Помесячная выручка совпадает с датами и суммами транзакций", failures)
            check(shares_ok, "Доля каждого месяца рассчитана от годовой выручки; сумма долей — 100%", failures)
            check(avg_ok, "Среднее значение указано для всех 12 месяцев", failures)
            charts = month_sheet._charts
            chart_ok = False
            if len(charts) == 1 and isinstance(charts[0], LineChart) and len(charts[0].series) == 2:
                chart = charts[0]
                categories = [series.cat.numRef.f if series.cat and series.cat.numRef else None for series in chart.series]
                revenues = [series.val.numRef.f if series.val and series.val.numRef else None for series in chart.series]
                expected_categories = "'Месячная выручка'!$A$2:$A$13"
                chart_ok = categories == [expected_categories, expected_categories] and revenues == [
                    "'Месячная выручка'!$B$2:$B$13",
                    "'Месячная выручка'!$F$2:$F$13",
                ]
            check(chart_ok, "Линейный график использует месяцы по порядку, выручку и серию среднего", failures)

            summary = analysis_wb["Итоги"]
            metrics = {summary.cell(row, 1).value: summary.cell(row, 2).value for row in range(2, summary.max_row + 1)}
            metrics_ok = True
            for key, expected in EXPECTED.items():
                actual = metrics.get(key)
                if isinstance(expected, int):
                    metrics_ok &= dec(actual) == Decimal(expected)
                else:
                    metrics_ok &= actual == expected
            metrics_ok &= month_names[max_index] == "Ноябрь" and monthly[max_index] == Decimal("1794540")
            metrics_ok &= month_names[min_index] == "Сентябрь" and monthly[min_index] == Decimal("127949")
            check(metrics_ok, "Итоговые контрольные показатели и журнал совпадают с набором задания", failures)
            log_sheet = analysis_wb["Журнал очистки"]
            check(log_sheet.max_row == 48, "Журнал содержит заголовок и все 47 исправлений", failures)


        outlier = next((row for row in clean_rows if str(row[0]) == "TXN-0001"), None)
        check(outlier is not None and dec(outlier[5]) == Decimal("897270"), "TXN-0001 сохранён как отмеченный выброс без выдуманной подстановки", failures)
    finally:
        source_wb.close()
        clean_wb.close()
        analysis_wb.close()

    print("\nЛокальная предварительная проверка завершена.")
    print("Это не официальный Analytics_autotest и не выдаёт контрольный код 7.1.")
    if failures:
        print(f"Провалено проверок: {len(failures)}")
        return 1
    print("Все проверки, которые можно вывести из приложенного ТЗ, пройдены.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
