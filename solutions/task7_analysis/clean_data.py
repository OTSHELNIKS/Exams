from __future__ import annotations

import argparse
import hashlib
import re
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = Path(__file__).resolve().parent
REPO_ROOT = HERE.parents[1]
HEADERS = ["transaction_id", "date", "product", "category", "quantity", "price", "shop", "manager"]
MONTHS_RU = [
    "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
    "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь",
]
PRODUCT_NAMES = {
    "утюг philips": "Утюг Philips",
    "микроволновая печь midea": "Микроволновая печь Midea",
    "чайник kitfort": "Чайник Kitfort",
    "пылесос dyson": "Пылесос Dyson",
    "пылесос dison": "Пылесос Dyson",
    "стиральная машина bosch": "Стиральная машина Bosch",
    "телевизор samsung": "Телевизор Samsung",
    "телевизор samsung 4k": "Телевизор Samsung 4K",
    "холодильник lg": "Холодильник LG",
}
MANAGERS = {
    "иванов и.и.": "Иванов И.И.",
    "смирнова е.в.": "Смирнова Е.В.",
    "сидоров в.к.": "Сидоров В.К.",
    "петрова а.с.": "Петрова А.С.",
}


def parse_numeric(value: Any, field_name: str) -> Decimal:
    if value is None:
        raise ValueError(f"Пустое значение в поле {field_name}")
    if isinstance(value, (int, float, Decimal)):
        try:
            parsed = Decimal(str(value))
        except InvalidOperation as error:
            raise ValueError(f"Некорректное число {value!r} в поле {field_name}") from error
    else:
        text = str(value).replace("\u00a0", " ").strip()

        text = re.sub(r"[^0-9,.-]", "", text)
        if "," in text and "." in text:


            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", ".")
        try:
            parsed = Decimal(text)
        except InvalidOperation as error:
            raise ValueError(f"Некорректное число {value!r} в поле {field_name}") from error
    if not parsed.is_finite():
        raise ValueError(f"Не конечное число в поле {field_name}: {value!r}")
    return parsed


def normalize_transaction_id(value: Any) -> str:
    text = str(value).strip().upper()
    match = re.fullmatch(r"TXN-?(\d+)", text)
    if not match:
        raise ValueError(f"Неизвестный формат transaction_id: {value!r}")
    return f"TXN-{int(match.group(1)):04d}"


def normalize_product(value: Any) -> str:
    text = " ".join(str(value or "").split())
    key = text.casefold()
    if key in PRODUCT_NAMES:
        return PRODUCT_NAMES[key]
    if not text:
        raise ValueError("Пустое название товара")
    return text


def infer_category(product: str) -> str:
    normalized = product.casefold()
    if normalized.startswith("пылесос"):
        return "Пылесосы"
    if normalized.startswith("телевизор"):
        return "Телевизоры"
    if normalized.startswith("стиральная машина"):
        return "Стиральные машины"
    if normalized.startswith("холодильник"):
        return "Холодильники"
    if normalized.startswith(("утюг", "чайник", "микроволновая печь")):
        return "Мелкая бытовая техника"
    raise ValueError(f"Для товара не настроена категория: {product}")


def normalize_manager(value: Any) -> str:
    text = " ".join(str(value or "").split())
    return MANAGERS.get(text.casefold(), text)


def as_excel_number(value: Decimal) -> int | float:
    return int(value) if value == value.to_integral_value() else float(value)


def load_and_clean(source: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]], int]:
    workbook = load_workbook(source, data_only=True, read_only=True)
    sheet = workbook.active
    all_rows = list(sheet.iter_rows(values_only=True))
    if not all_rows:
        raise ValueError("Исходная книга пуста")
    source_headers = list(all_rows[0][:8])
    if source_headers != HEADERS:
        raise ValueError(f"Ожидались колонки {HEADERS}, получены {source_headers}")

    raw_records: list[tuple[Any, ...]] = []
    for row in all_rows[1:]:

        first_eight = tuple(row[:8])
        if all(value is None for value in first_eight):
            continue
        if len(first_eight) != 8:
            raise ValueError(f"Строка не содержит 8 полей: {first_eight!r}")
        raw_records.append(first_eight)
    exact_duplicates = len(raw_records) - len(set(raw_records))

    cleaned: list[dict[str, Any]] = []
    change_log: list[dict[str, Any]] = []
    for excel_row, row in enumerate(raw_records, start=2):
        old = dict(zip(HEADERS, row))
        tx_id = normalize_transaction_id(old["transaction_id"])
        raw_date = old["date"]
        if isinstance(raw_date, datetime):
            trans_date = raw_date
        elif isinstance(raw_date, date):
            trans_date = datetime.combine(raw_date, datetime.min.time())
        else:
            trans_date = datetime.fromisoformat(str(raw_date).strip())
        if trans_date.year != 2026:
            raise ValueError(f"Ожидалась дата 2026 года, строка {excel_row}: {trans_date}")

        product = normalize_product(old["product"])
        category = infer_category(product)
        quantity_decimal = parse_numeric(old["quantity"], "quantity")
        if quantity_decimal <= 0 or quantity_decimal != quantity_decimal.to_integral_value():
            raise ValueError(f"Количество должно быть положительным целым числом, строка {excel_row}")
        price_decimal = parse_numeric(old["price"], "price")
        if price_decimal < 0:
            raise ValueError(f"Цена не может быть отрицательной, строка {excel_row}")
        shop = " ".join(str(old["shop"] or "").split())
        manager = normalize_manager(old["manager"])

        clean = {
            "transaction_id": tx_id,
            "date": trans_date,
            "product": product,
            "category": category,
            "quantity": int(quantity_decimal),
            "price": as_excel_number(price_decimal),
            "shop": shop,
            "manager": manager,
        }
        cleaned.append(clean)
        for field in HEADERS:
            previous = old[field]
            current = clean[field]
            if field == "date" and isinstance(previous, (datetime, date)):
                unchanged = previous.date() == current.date() if isinstance(previous, datetime) else previous == current.date()
            elif field in ("quantity", "price"):


                unchanged = isinstance(previous, (int, float, Decimal)) and not isinstance(previous, bool) \
                    and Decimal(str(previous)) == Decimal(str(current))
            else:
                unchanged = previous == current
            if not unchanged:
                change_log.append({
                    "excel_row": excel_row,
                    "transaction_id": tx_id,
                    "field": field,
                    "old_value": "" if previous is None else str(previous),
                    "new_value": "" if current is None else str(current),
                })

    ids = [row["transaction_id"] for row in cleaned]
    if len(ids) != len(set(ids)):
        raise ValueError("После стандартизации обнаружены повторяющиеся transaction_id")
    workbook.close()
    return cleaned, change_log, exact_duplicates


def style_table(sheet, header_row: int = 1) -> None:
    navy = "172B4D"
    for cell in sheet[header_row]:
        if cell.value is not None:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[header_row].height = 30
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.sheet_view.showGridLines = False


def add_excel_table(sheet, name: str, ref: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False,
        showLastColumn=False, showRowStripes=True, showColumnStripes=False,
    )
    sheet.add_table(table)


def write_clean_workbook(output: Path, cleaned: list[dict[str, Any]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(HEADERS)
    for row in cleaned:
        sheet.append([row[field] for field in HEADERS])
    style_table(sheet)
    widths = [17, 14, 34, 30, 13, 14, 22, 22]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 2).number_format = "DD.MM.YYYY"
        sheet.cell(row, 5).number_format = "0"
        sheet.cell(row, 6).number_format = '#,##0.00;[Red]-#,##0.00'
    if cleaned:
        add_excel_table(sheet, "CleanTransactions", f"A1:H{sheet.max_row}")
    workbook.properties.title = "Очищенный набор транзакций розничной сети за 2026 год"
    workbook.properties.subject = "Задание 7 — предобработка без удаления строк"
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def write_analysis_workbook(
    output: Path,
    cleaned: list[dict[str, Any]],
    change_log: list[dict[str, Any]],
    exact_duplicates: int,
) -> dict[str, Any]:
    monthly = defaultdict(Decimal)
    data_revenues: list[Decimal] = []
    for row in cleaned:
        revenue = Decimal(str(row["quantity"])) * Decimal(str(row["price"]))
        data_revenues.append(revenue)
        monthly[row["date"].month] += revenue
    total = sum(monthly.values(), Decimal("0"))
    average = total / Decimal(12)
    min_month = min(range(1, 13), key=lambda month: monthly[month])
    max_month = max(range(1, 13), key=lambda month: monthly[month])

    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "Данные"
    data_sheet.append(HEADERS + ["revenue"])
    for row, revenue in zip(cleaned, data_revenues):
        data_sheet.append([row[field] for field in HEADERS] + [as_excel_number(revenue)])
    style_table(data_sheet)
    for idx, width in enumerate([17, 14, 34, 30, 13, 14, 22, 22, 18], start=1):
        data_sheet.column_dimensions[chr(64 + idx)].width = width
    for row_num in range(2, data_sheet.max_row + 1):
        data_sheet.cell(row_num, 2).number_format = "DD.MM.YYYY"
        for col in (5, 6, 9):
            data_sheet.cell(row_num, col).number_format = '#,##0.00;[Red]-#,##0.00'
    add_excel_table(data_sheet, "AnalysisTransactions", f"A1:I{data_sheet.max_row}")

    month_sheet = workbook.create_sheet("Месячная выручка")
    month_headers = [
        "Месяц", "Выручка, руб.", "Доля годовой выручки, %",
        "Изменение к предыдущему месяцу, руб.", "Изменение к предыдущему месяцу, %",
        "Среднее за месяц, руб.",
    ]
    month_sheet.append(month_headers)
    for month in range(1, 13):
        value = monthly[month]
        previous = monthly[month - 1] if month > 1 else None
        change = value - previous if previous is not None else None
        change_pct = (change / previous * 100) if previous not in (None, Decimal("0")) else None
        month_sheet.append([
            MONTHS_RU[month - 1],
            as_excel_number(value),
            float(value / total * 100) if total else 0,
            as_excel_number(change) if change is not None else None,
            float(change_pct) if change_pct is not None else None,
            as_excel_number(average),
        ])
    style_table(month_sheet)
    for idx, width in enumerate([17, 20, 24, 36, 35, 24], start=1):
        month_sheet.column_dimensions[chr(64 + idx)].width = width
    for row_num in range(2, 14):
        for col in (2, 4, 6):
            month_sheet.cell(row_num, col).number_format = '#,##0.00;[Red]-#,##0.00'
        for col in (3, 5):
            month_sheet.cell(row_num, col).number_format = '0.00'
    add_excel_table(month_sheet, "MonthlyRevenue", "A1:F13")

    chart = LineChart()
    chart.title = "Динамика месячной выручки за 2026 год"
    chart.style = 13
    chart.y_axis.title = "Выручка, руб."
    chart.x_axis.title = "Месяц"
    chart.height = 10
    chart.width = 21
    chart.legend.position = "b"
    chart.y_axis.numFmt = '#,##0'
    revenue_data = Reference(month_sheet, min_col=2, max_col=2, min_row=1, max_row=13)
    average_data = Reference(month_sheet, min_col=6, max_col=6, min_row=1, max_row=13)
    categories = Reference(month_sheet, min_col=1, min_row=2, max_row=13)
    chart.add_data(revenue_data, titles_from_data=True)
    chart.add_data(average_data, titles_from_data=True)
    chart.set_categories(categories)
    chart.series[0].graphicalProperties.line.solidFill = "4058D6"
    chart.series[0].graphicalProperties.line.width = 28575
    chart.series[1].graphicalProperties.line.solidFill = "E0565B"
    chart.series[1].graphicalProperties.line.width = 19050
    chart.series[1].graphicalProperties.line.prstDash = "dash"
    month_sheet.add_chart(chart, "H2")
    month_sheet.freeze_panes = "A2"

    summary = workbook.create_sheet("Итоги")
    summary.append(["Показатель", "Значение"])
    metrics = [
        ("Строк в исходном наборе", len(cleaned)),
        ("Полных дубликатов", exact_duplicates),
        ("Строк удалено", 0),
        ("Годовая выручка, руб.", as_excel_number(total)),
        ("Средняя месячная выручка, руб.", as_excel_number(average)),
        ("Месяц максимальной выручки", MONTHS_RU[max_month - 1]),
        ("Максимальная месячная выручка, руб.", as_excel_number(monthly[max_month])),
        ("Месяц минимальной выручки", MONTHS_RU[min_month - 1]),
        ("Минимальная месячная выручка, руб.", as_excel_number(monthly[min_month])),
        ("Размах месячной выручки, руб.", as_excel_number(monthly[max_month] - monthly[min_month])),
        ("Строк с зафиксированными исправлениями", len({x["excel_row"] for x in change_log})),
        ("Количество исправлений по полям", len(change_log)),
    ]
    for item in metrics:
        summary.append(list(item))
    summary.append(["Интерпретация", "Выручка волатильна: максимум — ноябрь, минимум — сентябрь; устойчивого ежемесячного роста нет."])
    summary.append(["Проверить по первичному источнику", "TXN-0001: 897270 руб./шт. за утюг Philips — явный ценовой выброс. В clean-файле сохранён как получен: в приложенных условиях нет подтверждённого исправленного значения."])
    style_table(summary)
    summary.column_dimensions["A"].width = 43
    summary.column_dimensions["B"].width = 92
    summary.freeze_panes = "A2"

    log_sheet = workbook.create_sheet("Журнал очистки")
    log_headers = ["Строка Excel", "transaction_id", "Поле", "Исходное значение", "Очищенное значение"]
    log_sheet.append(log_headers)
    for item in change_log:
        log_sheet.append([item["excel_row"], item["transaction_id"], item["field"], item["old_value"], item["new_value"]])
    style_table(log_sheet)
    for col, width in {"A": 15, "B": 18, "C": 18, "D": 38, "E": 38}.items():
        log_sheet.column_dimensions[col].width = width
    if change_log:
        add_excel_table(log_sheet, "CleaningLog", f"A1:E{log_sheet.max_row}")

    workbook.properties.title = "Анализ продаж за 2026 год"
    workbook.properties.subject = "Месячная выручка, среднее значение и временной ряд"
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return {
        "monthly": monthly,
        "total": total,
        "average": average,
        "min_month": min_month,
        "max_month": max_month,
        "change_log": change_log,
        "duplicates": exact_duplicates,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description='Очистка транзакций 2026 года и формирование итоговых книг.')
    parser.add_argument("--input", type=Path, default=REPO_ROOT / "dataset_2026.xlsx")
    parser.add_argument("--clean-output", type=Path, default=HERE / "clean_dataset_2026.xlsx")
    parser.add_argument("--analysis-output", type=Path, default=HERE / "analysis_2026.xlsx")
    args = parser.parse_args()

    cleaned, change_log, duplicate_count = load_and_clean(args.input)
    write_clean_workbook(args.clean_output, cleaned)
    result = write_analysis_workbook(args.analysis_output, cleaned, change_log, duplicate_count)
    digest = hashlib.sha256(args.clean_output.read_bytes()).hexdigest()
    print(f"Строк в очищенном файле: {len(cleaned)}; удалено: 0; полных дубликатов: {duplicate_count}")
    print(f"Исправлений в журнале: {len(change_log)}; затронуто строк: {len({item['excel_row'] for item in change_log})}")
    print(f"Годовая выручка: {result['total']}; средняя за месяц: {result['average']}")
    print(f"Максимум: {MONTHS_RU[result['max_month'] - 1]} — {result['monthly'][result['max_month']]}")
    print(f"Минимум: {MONTHS_RU[result['min_month'] - 1]} — {result['monthly'][result['min_month']]}")
    print(f"SHA-256 файла очистки (фингерпринт, не контрольный код приложения): {digest}")
    print(f"Очищенный набор: {args.clean_output}")
    print(f"Расчёты и график: {args.analysis_output}")


if __name__ == "__main__":
    main()
